from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from services.placeholders import replace_tokens


def _replace_in_text_frame(text_frame, values: dict[str, str], keep_unanswered: bool) -> None:
    for paragraph in text_frame.paragraphs:
        if not paragraph.runs:
            paragraph.text = replace_tokens(paragraph.text, values, keep_unanswered)
            continue

        full = "".join(run.text for run in paragraph.runs)
        replaced = replace_tokens(full, values, keep_unanswered)
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""


def fill_pptx(template_path: str, output_path: str, values: dict[str, str], keep_unanswered: bool = False) -> str:
    prs = Presentation(template_path)

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame is not None:
                _replace_in_text_frame(shape.text_frame, values, keep_unanswered)
            if hasattr(shape, "table") and shape.table is not None:
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = replace_tokens(cell.text or "", values, keep_unanswered)

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = replace_tokens(notes_tf.text or "", values, keep_unanswered)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return str(out)
