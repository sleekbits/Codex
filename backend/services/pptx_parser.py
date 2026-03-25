from __future__ import annotations

from pptx import Presentation

from models.job import PlaceholderMeta
from services.placeholders import find_placeholders, infer_type


def parse_pptx_placeholders(path: str, include_notes: bool = False) -> list[PlaceholderMeta]:
    prs = Presentation(path)
    found: dict[tuple[str, int, str | None], PlaceholderMeta] = {}

    for sidx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            shape_name = getattr(shape, "name", None)
            if hasattr(shape, "text_frame") and shape.text_frame is not None:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
                    for key, fmt in find_placeholders(text):
                        meta = PlaceholderMeta(
                            key=key,
                            format=fmt,
                            slideIndex=sidx,
                            shapeName=shape_name,
                            context=text[:160],
                            suggestedType=infer_type(key, text),
                        )
                        found[(key, sidx, shape_name)] = meta
            if hasattr(shape, "table") and shape.table is not None:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text or ""
                        for key, fmt in find_placeholders(text):
                            meta = PlaceholderMeta(
                                key=key,
                                format=fmt,
                                slideIndex=sidx,
                                shapeName=shape_name,
                                context=text[:160],
                                suggestedType=infer_type(key, text),
                            )
                            found[(key, sidx, shape_name)] = meta

        if include_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            for key, fmt in find_placeholders(notes):
                meta = PlaceholderMeta(
                    key=key,
                    format=fmt,
                    slideIndex=sidx,
                    shapeName="notes",
                    context=notes[:160],
                    suggestedType=infer_type(key, notes),
                )
                found[(key, sidx, "notes")] = meta

    return sorted(found.values(), key=lambda p: (p.slideIndex, p.key))
