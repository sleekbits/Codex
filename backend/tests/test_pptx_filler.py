from pathlib import Path

from pptx import Presentation

from services.pptx_filler import fill_pptx


def _make_template(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(left=0, top=0, width=3000000, height=1000000)
    tf = box.text_frame
    tf.text = "Hello {{projectTitle}}"

    table = slide.shapes.add_table(1, 1, 0, 1000000, 4000000, 1000000).table
    table.cell(0, 0).text = "Budget {{budget}}"
    prs.save(path)


def test_fill_replaces_text_and_table(tmp_path: Path):
    template = tmp_path / "template.pptx"
    output = tmp_path / "output.pptx"
    _make_template(template)

    fill_pptx(str(template), str(output), {"projectTitle": "Alpha", "budget": "120000"})

    prs = Presentation(str(output))
    texts = []
    for shape in prs.slides[0].shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            texts.append(shape.text_frame.text)
        if hasattr(shape, "table") and shape.table:
            texts.append(shape.table.cell(0, 0).text)

    joined = "\n".join(texts)
    assert "Hello Alpha" in joined
    assert "Budget 120000" in joined
