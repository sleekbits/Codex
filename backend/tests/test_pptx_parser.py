from pathlib import Path

from pptx import Presentation

from services.pptx_parser import parse_pptx_placeholders


def _create_fixture(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(left=0, top=0, width=3000000, height=1000000)
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.add_run().text = "Award {{pro"
    p.add_run().text = "jectTitle}}"
    p2 = tf.add_paragraph()
    p2.text = "Client [[clientName]]"

    table = slide.shapes.add_table(2, 2, 0, 1000000, 4000000, 2000000).table
    table.cell(0, 0).text = "Budget <<budget>>"
    prs.save(path)


def test_extracts_placeholders(tmp_path: Path):
    pptx_path = tmp_path / "sample.pptx"
    _create_fixture(pptx_path)

    data = parse_pptx_placeholders(str(pptx_path))
    keys = {p.key for p in data}

    assert "projectTitle" in keys
    assert "clientName" in keys
    assert "budget" in keys
